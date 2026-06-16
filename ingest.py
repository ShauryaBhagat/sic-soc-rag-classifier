# ingest.py
# Loads SIC 2007 index (ONS Excel format) into ChromaDB vector store.
# Each activity description becomes a searchable document pointing to its SIC code.
# Run this once before using retrieve.py or classify.py.

import os
from dotenv import load_dotenv
from langchain_openai import OpenAIEmbeddings
from langchain_chroma import Chroma
from langchain_core.documents import Document
import openpyxl

load_dotenv()

DATA_PATH = "data/SIC_2007_index.xlsx"
CHROMA_PATH = "data/chroma_store"


def load_sic_codes(filepath: str) -> list[Document]:
    """
    Reads the ONS SIC 2007 Excel index and converts each row into a
    LangChain Document.

    The ONS file has 15,958 rows — multiple activity descriptions per
    SIC code. We treat each activity description as a separate searchable
    document, with the SIC code stored in metadata. This gives much richer
    retrieval than one-description-per-code because real job descriptions
    often match specific activities rather than broad category names.

    Column names in the file:
      - 'UK SIC 2007' : the 5-digit SIC code
      - 'Activity'    : a specific activity description for that code
    """
    print(f"Loading SIC data from: {filepath}")

    wb = openpyxl.load_workbook(filepath, read_only=True)
    ws = wb.active

    rows = list(ws.iter_rows(values_only=True))
    header = rows[0]
    print(f"Columns found: {header}")
    # Expected: ('UK SIC 2007', 'Activity')

    documents = []
    skipped = 0

    for row in rows[1:]:  # skip header row
        sic_code = str(row[0]).strip() if row[0] else ""
        activity = str(row[1]).strip() if row[1] else ""

        if not sic_code or not activity or sic_code == "None":
            skipped += 1
            continue

        doc = Document(
            page_content=activity,   # this is what gets embedded and searched
            metadata={
                "sic_code": sic_code,
                "activity": activity
            }
        )
        documents.append(doc)

    print(f"Loaded {len(documents)} activity descriptions across SIC codes")
    print(f"Skipped {skipped} empty/malformed rows")
    return documents


def build_vector_store(documents: list[Document]) -> Chroma:
    """
    Embeds all documents using OpenAI text-embedding-3-small and
    persists to ChromaDB.

    text-embedding-3-small costs ~$0.00002 per 1K tokens.
    15,958 short activity descriptions ≈ ~150K tokens total ≈ $0.003.
    This entire ingestion run costs less than half a penny.
    """
    print(f"\nEmbedding {len(documents)} documents...")
    print("This will take 1-2 minutes and cost ~$0.003 in API credits.")

    embeddings = OpenAIEmbeddings(model="text-embedding-3-small")

    # Chroma.from_documents embeds in batches and saves to disk
    vector_store = Chroma.from_documents(
        documents=documents,
        embedding=embeddings,
        persist_directory=CHROMA_PATH
    )

    count = vector_store._collection.count()
    print(f"\nVector store built at: {CHROMA_PATH}")
    print(f"Total documents stored: {count}")
    return vector_store


if __name__ == "__main__":
    if not os.getenv("OPENAI_API_KEY"):
        raise ValueError(
            "OPENAI_API_KEY not found. "
            "Make sure your .env file exists and contains the key."
        )

    if not os.path.exists(DATA_PATH):
        raise FileNotFoundError(
            f"Data file not found at {DATA_PATH}. "
            "Make sure you have placed the ONS Excel file in the data/ folder."
        )

    docs = load_sic_codes(DATA_PATH)
    store = build_vector_store(docs)
    print("\nIngestion complete. Run retrieve.py to test retrieval.")